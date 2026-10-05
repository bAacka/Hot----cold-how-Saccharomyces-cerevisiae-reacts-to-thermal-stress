#!/usr/bin/env python3

from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(
    "/"
)

TABLE6 = (
    ROOT
    / "34_phosphoatlas_external_validation"
    / "raw"
    / "PhosphoAtlas_SuppTable6.xlsx"
)

S1FILE = (
    ROOT
    / "00_raw"
    / "Kanshin_2015_Table_S1.xlsx"
)

DIRECT_TARGETS = (
    ROOT
    / "31_uliana_context_interaction"
    / "DIRECT_A_exact_targets.tsv"
)

NETWORKIN = (
    ROOT
    / "17_networkin_temporal_discovery"
    / "networkin_temporal_candidate_membership.tsv"
)

OUT = (
    ROOT
    / "36_phosphoatlas_temperature_sensitivity"
)

CACHE = (
    ROOT
    / "34_phosphoatlas_external_validation"
    / "processed"
    / "PhosphoAtlas_temperature_conditions.tsv.gz"
)

OUT.mkdir(
    parents=True,
    exist_ok=True,
)

CONDITIONS = [
    "HS37",
    "HS42",
    "HS48",
    "CS18",
    "CS23",
]

CONTRASTS = [
    (
        "HS37_minus_CS18",
        "HS37",
        "CS18",
    ),
    (
        "HS48_minus_CS18",
        "HS48",
        "CS18",
    ),
    (
        "HS42_minus_CS23",
        "HS42",
        "CS23",
    ),
]

ORTH_LABEL = (
    "TPK2_TPK1_TPK3_group"
    "__ORTHOGONAL_NO_CURATED_AB"
)

N_PERM = 100000
SEED = 20260915


def clean(x):
    if pd.isna(x):
        return ""
    return str(x).strip()


def norm_id(x):
    if pd.isna(x):
        return ""

    s = str(x).strip()

    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
    except Exception:
        pass

    return s


def parse_sites(x):
    return re.findall(
        r"[STY]\d+",
        clean(x),
    )


def empirical_lower(
    obs,
    null,
):
    return float(
        (
            1
            + np.sum(
                null <= obs
            )
        )
        / (
            len(null)
            + 1
        )
    )


def bh(values):
    p = np.asarray(
        values,
        dtype=float,
    )

    out = np.full(
        len(p),
        np.nan,
    )

    good = np.isfinite(p)

    if not good.any():
        return out

    pv = p[good]

    order = np.argsort(
        pv
    )

    ranked = pv[
        order
    ]

    q = (
        ranked
        * len(ranked)
        / np.arange(
            1,
            len(ranked) + 1,
        )
    )

    q = np.minimum.accumulate(
        q[::-1]
    )[::-1]

    q = np.minimum(
        q,
        1.0,
    )

    restored = np.empty_like(
        q
    )

    restored[
        order
    ] = q

    out[
        good
    ] = restored

    return out


                                                              
                                                 
                                                              

if CACHE.exists():

    temp = pd.read_csv(
        CACHE,
        sep="\t",
        compression="gzip",
    )

else:

    print(
        "Streaming Table S6 for temperature treatments..."
    )

    wb = load_workbook(
        TABLE6,
        read_only=True,
        data_only=True,
    )

    ws = wb[
        "p_site_diff_reg"
    ]

    rows = ws.iter_rows(
        values_only=True
    )

    header = next(
        rows
    )

    col = {
        str(v): i
        for i, v
        in enumerate(
            header
        )
    }

    wanted = [
        "reference",
        "systematic_name",
        "gene",
        "p_site",
        "p_residue",
        "p_position",
        "treatment_id",
        "fc_log2",
        "p_value",
        "adj_p_value",
    ]

    out = []

    for row in rows:

        treatment = row[
            col[
                "treatment_id"
            ]
        ]

        if treatment not in CONDITIONS:
            continue

        out.append({
            x:
                row[
                    col[x]
                ]

            for x in wanted
        })

    temp = pd.DataFrame(
        out
    )

    CACHE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temp.to_csv(
        CACHE,
        sep="\t",
        index=False,
        compression="gzip",
    )


print()
print("TEMPERATURE ROW COUNTS")
print(
    temp[
        "treatment_id"
    ].value_counts()
)


                                                              
                                               
                                                              

temp["gene"] = (
    temp["gene"]
    .where(
        temp["gene"].notna(),
        "",
    )
    .astype(str)
    .str.strip()
    .str.upper()
)

temp["p_residue"] = (
    temp[
        "p_residue"
    ]
    .astype(str)
    .str.strip()
)

temp["p_position"] = pd.to_numeric(
    temp[
        "p_position"
    ],
    errors="coerce",
)

temp["fc_log2"] = pd.to_numeric(
    temp[
        "fc_log2"
    ],
    errors="coerce",
)

temp["adj_p_value"] = pd.to_numeric(
    temp[
        "adj_p_value"
    ],
    errors="coerce",
)


temp = temp[
    (temp["gene"] != "")
    & temp[
        "p_residue"
    ].isin(
        [
            "S",
            "T",
            "Y",
        ]
    )
    & temp[
        "p_position"
    ].notna()
    & temp[
        "fc_log2"
    ].notna()
].copy()


temp[
    "p_position"
] = temp[
    "p_position"
].astype(int)


temp = (
    temp
    .groupby(
        [
            "gene",
            "p_residue",
            "p_position",
            "treatment_id",
        ],
        as_index=False,
    )
    .agg(
        fc_log2=(
            "fc_log2",
            "median",
        ),

        adj_p_value=(
            "adj_p_value",
            "min",
        ),
    )
)


                                                              
                                
                                                              

s1 = pd.read_excel(
    S1FILE,
    sheet_name="Phosphopeptides",
)


gene_col = (
    "Gene"
    if "Gene" in s1.columns
    else "Standard Name"
)


s1["s1_id"] = (
    s1["id"]
    .map(
        norm_id
    )
)


s1["gene"] = (
    s1[
        gene_col
    ]
    .where(
        s1[
            gene_col
        ].notna(),
        "",
    )
    .astype(str)
    .str.strip()
    .str.upper()
)


s1["site_id"] = (
    s1["gene"]
    + "_"
    + s1[
        "pSites"
    ].astype(str)
)


atom_rows = []


for _, r in s1.iterrows():

    if not r["gene"]:
        continue

    for site in parse_sites(
        r[
            "pSites"
        ]
    ):

        m = re.fullmatch(
            r"([STY])(\d+)",
            site,
        )

        if not m:
            continue

        atom_rows.append({
            "s1_id":
                r[
                    "s1_id"
                ],

            "gene":
                r[
                    "gene"
                ],

            "pSites":
                clean(
                    r[
                        "pSites"
                    ]
                ),

            "site_id":
                r[
                    "site_id"
                ],

            "atom_site":
                site,

            "p_residue":
                m.group(1),

            "p_position":
                int(
                    m.group(2)
                ),
        })


atoms = pd.DataFrame(
    atom_rows
)


                                                              
                      
                                                              

direct = pd.read_csv(
    DIRECT_TARGETS,
    sep="\t",
    dtype={
        "s1_id":
            str,
    },
)

direct["s1_id"] = (
    direct[
        "s1_id"
    ].map(
        norm_id
    )
)

direct_ids = set(
    direct[
        "s1_id"
    ]
)


mem = pd.read_csv(
    NETWORKIN,
    sep="\t",
)


orth = mem[
    mem[
        "analysis_label"
    ] == ORTH_LABEL
].copy()


orth["gene"] = (
    orth["gene"]
    .astype(str)
    .str.strip()
    .str.upper()
)


orth = orth.drop_duplicates(
    [
        "gene",
        "site_id",
    ]
)


orth_keys = set(
    zip(
        orth["gene"],
        orth["site_id"],
    )
)


orth_s1 = s1[
    s1.apply(
        lambda r:
            (
                r[
                    "gene"
                ],
                r[
                    "site_id"
                ],
            )
            in orth_keys,
        axis=1,
    )
].copy()


orth_ids = set(
    orth_s1[
        "s1_id"
    ]
)


if len(
    direct_ids
) != 10:

    raise RuntimeError(
        "Direct-A membership changed"
    )


if len(
    orth_ids
) != 20:

    raise RuntimeError(
        "Orthogonal PKA membership changed"
    )


MODULES = {
    "DIRECT_A_PKA":
        direct_ids,

    "ORTHOGONAL_NETWORKIN_PKA":
        orth_ids,
}


                                                              
                       
                                                              

def build_contrast(
    label,
    heat_id,
    cold_id,
):

    h = temp[
        temp[
            "treatment_id"
        ] == heat_id
    ].copy()

    c = temp[
        temp[
            "treatment_id"
        ] == cold_id
    ].copy()


    h = h.rename(
        columns={
            "fc_log2":
                "heat_fc",

            "adj_p_value":
                "heat_q",
        }
    ).drop(
        columns=[
            "treatment_id"
        ]
    )


    c = c.rename(
        columns={
            "fc_log2":
                "cold_fc",

            "adj_p_value":
                "cold_q",
        }
    ).drop(
        columns=[
            "treatment_id"
        ]
    )


    ext = h.merge(
        c,
        on=[
            "gene",
            "p_residue",
            "p_position",
        ],
        how="inner",
        validate="1:1",
    )


    ext[
        "external_delta"
    ] = (
        ext[
            "heat_fc"
        ]
        -
        ext[
            "cold_fc"
        ]
    )


    ma = atoms.merge(
        ext,
        on=[
            "gene",
            "p_residue",
            "p_position",
        ],
        how="left",
        validate="m:1",
    )


    ma[
        "matched"
    ] = ma[
        "external_delta"
    ].notna()


    group_rows = []


    for (
        s1_id,
        gene,
        psites,
        site_id,
    ), g in ma.groupby(
        [
            "s1_id",
            "gene",
            "pSites",
            "site_id",
        ],
        sort=False,
    ):

        gm = g[
            g[
                "matched"
            ]
        ]

        if len(
            gm
        ) == 0:
            continue

        group_rows.append({
            "s1_id":
                s1_id,

            "gene":
                gene,

            "pSites":
                psites,

            "site_id":
                site_id,

            "n_atoms":
                len(
                    g
                ),

            "n_matched_atoms":
                len(
                    gm
                ),

            "external_delta":
                float(
                    np.median(
                        gm[
                            "external_delta"
                        ]
                    )
                ),

            "heat_fc":
                float(
                    np.median(
                        gm[
                            "heat_fc"
                        ]
                    )
                ),

            "cold_fc":
                float(
                    np.median(
                        gm[
                            "cold_fc"
                        ]
                    )
                ),
        })


    groups = pd.DataFrame(
        group_rows
    )


    groups.to_csv(
        OUT
        / f"{label}_all_matched_groups.tsv.gz",
        sep="\t",
        index=False,
        compression="gzip",
    )


    return groups


                                                              
                                     
                                                              

def module_test(
    groups,
    module_name,
    target_ids,
    contrast_label,
    seed,
):

    target = groups[
        groups[
            "s1_id"
        ].isin(
            target_ids
        )
    ].copy()


    target_genes = set(
        target[
            "gene"
        ]
    )


    bg = groups[
        ~groups[
            "gene"
        ].isin(
            target_genes
        )
        &
        ~groups[
            "s1_id"
        ].isin(
            target_ids
        )
    ].copy()


    tg = (
        target
        .groupby(
            "gene",
            as_index=False,
        )
        .agg(
            n_groups=(
                "s1_id",
                "nunique",
            ),

            value=(
                "external_delta",
                "median",
            ),
        )
    )


    if len(
        tg
    ) < 3:

        raise RuntimeError(
            f"{contrast_label} / {module_name}: "
            f"only {len(tg)} target genes"
        )


    mults = sorted(
        tg[
            "n_groups"
        ].astype(int).tolist(),
        reverse=True,
    )


    obs = float(
        tg[
            "value"
        ].mean()
    )


    obs_median = float(
        tg[
            "value"
        ].median()
    )


    n_negative = int(
        (
            tg[
                "value"
            ] < 0
        ).sum()
    )


    by_gene = {
        gene:
            g[
                "external_delta"
            ].to_numpy(
                dtype=float
            )

        for gene, g
        in bg.groupby(
            "gene"
        )
    }


    eligible = {
        m: np.array(
            [
                gene
                for gene, vals
                in by_gene.items()
                if len(
                    vals
                ) >= m
            ],
            dtype=object,
        )

        for m in set(
            mults
        )
    }


    rng = np.random.default_rng(
        seed
    )


    null_mean = np.empty(
        N_PERM
    )

    null_median = np.empty(
        N_PERM
    )

    null_negative = np.empty(
        N_PERM,
        dtype=int,
    )


    for i in range(
        N_PERM
    ):

        chosen = set()
        vals_out = []

        for m in mults:

            pool = eligible[
                m
            ]

            while True:

                gene = pool[
                    rng.integers(
                        len(
                            pool
                        )
                    )
                ]

                if gene not in chosen:
                    break

            chosen.add(
                gene
            )

            vals = by_gene[
                gene
            ]

            idx = rng.choice(
                len(
                    vals
                ),
                size=m,
                replace=False,
            )

            vals_out.append(
                float(
                    np.median(
                        vals[
                            idx
                        ]
                    )
                )
            )


        vals_out = np.asarray(
            vals_out
        )


        null_mean[i] = np.mean(
            vals_out
        )

        null_median[i] = np.median(
            vals_out
        )

        null_negative[i] = np.sum(
            vals_out < 0
        )


    mu = float(
        np.mean(
            null_mean
        )
    )

    sd = float(
        np.std(
            null_mean,
            ddof=1,
        )
    )


    p_mean = empirical_lower(
        obs,
        null_mean,
    )


    p_median = empirical_lower(
        obs_median,
        null_median,
    )


    p_negative = float(
        (
            1
            + np.sum(
                null_negative
                >= n_negative
            )
        )
        / (
            N_PERM
            + 1
        )
    )


    target.to_csv(
        OUT
        / (
            f"{contrast_label}"
            f"__{module_name}"
            "_matched_targets.tsv"
        ),
        sep="\t",
        index=False,
    )


    return {
        "contrast":
            contrast_label,

        "module":
            module_name,

        "predefined_groups":
            len(
                target_ids
            ),

        "matched_groups":
            len(
                target
            ),

        "matched_genes":
            len(
                tg
            ),

        "multiplicity_pattern":
            ",".join(
                map(
                    str,
                    mults,
                )
            ),

        "observed_mean":
            obs,

        "observed_median":
            obs_median,

        "negative_genes":
            n_negative,

        "null_mean":
            mu,

        "null_sd":
            sd,

        "empirical_z":
            float(
                (
                    obs
                    - mu
                )
                / sd
            ),

        "mean_empirical_lower_p":
            p_mean,

        "median_empirical_lower_p":
            p_median,

        "negative_count_empirical_p":
            p_negative,
    }


                                                              
                                     
                                                              

results = []

seed_i = 0


for (
    contrast_label,
    heat_id,
    cold_id,
) in CONTRASTS:

    print()
    print(
        "Building",
        contrast_label,
    )

    groups = build_contrast(
        contrast_label,
        heat_id,
        cold_id,
    )

    print(
        "Matched S1 groups:",
        len(
            groups
        ),
    )


    for (
        module_name,
        ids,
    ) in MODULES.items():

        seed_i += 1

        result = module_test(
            groups,
            module_name,
            ids,
            contrast_label,
            SEED
            + seed_i,
        )

        results.append(
            result
        )


results = pd.DataFrame(
    results
)


                                                       
results[
    "q_across_six_sensitivity_tests"
] = bh(
    results[
        "mean_empirical_lower_p"
    ]
)


results.to_csv(
    OUT
    / "TEMPERATURE_SENSITIVITY_SUMMARY.tsv",
    sep="\t",
    index=False,
)


print()
print("=" * 110)
print("PHOSPHOATLAS TEMPERATURE SENSITIVITY")
print("=" * 110)

print(
    results.to_string(
        index=False
    )
)

print()
print(
    "Output:",
    OUT,
)
