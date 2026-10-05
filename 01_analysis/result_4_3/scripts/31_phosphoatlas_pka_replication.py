#!/usr/bin/env python3

from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import spearmanr, pearsonr, binomtest


ROOT = Path("/")

S1FILE = (
    ROOT
    / "00_raw"
    / "Kanshin_2015_Table_S1.xlsx"
)

TABLE6 = (
    ROOT
    / "34_phosphoatlas_external_validation"
    / "raw"
    / "PhosphoAtlas_SuppTable6.xlsx"
)

CACHE = (
    ROOT
    / "34_phosphoatlas_external_validation"
    / "processed"
    / "PhosphoAtlas_HS42_CS18.tsv.gz"
)

DIRECT_TARGETS = (
    ROOT
    / "31_uliana_context_interaction"
    / "DIRECT_A_exact_targets.tsv"
)

NETWORKIN = (
    ROOT
    / "17_networkin_temporal_discovery"
    / "networkin_temporal_candidate_membership.tsv"
)

OUT = (
    ROOT
    / "35_phosphoatlas_pka_replication"
)
OUT.mkdir(parents=True, exist_ok=True)

HEAT = "HS42"
COLD = "CS18"

ORTH_LABEL = (
    "TPK2_TPK1_TPK3_group"
    "__ORTHOGONAL_NO_CURATED_AB"
)

N_PERM = 100000
SEED = 20260915


                                                              
         
                                                              

def clean(x):
    if pd.isna(x):
        return ""
    return str(x).strip()


def norm_id(x):
    if pd.isna(x):
        return ""

    s = str(x).strip()

    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
    except Exception:
        pass

    return s


def parse_sites(x):
    return re.findall(
        r"[STY]\d+",
        clean(x),
    )


def bh(values):
    p = np.asarray(values, dtype=float)

    out = np.full(
        len(p),
        np.nan,
    )

    good = np.isfinite(p)

    if not good.any():
        return out

    pv = p[good]
    order = np.argsort(pv)
    ranked = pv[order]

    q = (
        ranked
        * len(ranked)
        / np.arange(
            1,
            len(ranked) + 1,
        )
    )

    q = np.minimum.accumulate(
        q[::-1]
    )[::-1]

    q = np.minimum(
        q,
        1.0,
    )

    restored = np.empty_like(q)
    restored[order] = q

    out[good] = restored

    return out


def empirical_lower(obs, null):
    null = np.asarray(
        null,
        dtype=float,
    )

    return float(
        (
            1
            + np.sum(
                null <= obs
            )
        )
        / (
            len(null)
            + 1
        )
    )


                                                              
                                                     
                                                              

if CACHE.exists():

    print(
        "Using cached PhosphoAtlas temperature subset:",
        CACHE,
    )

    atlas_raw = pd.read_csv(
        CACHE,
        sep="\t",
        compression="gzip",
    )

else:

    print(
        "Streaming PhosphoAtlas Table S6..."
    )

    wb = load_workbook(
        TABLE6,
        read_only=True,
        data_only=True,
    )

    ws = wb[
        "p_site_diff_reg"
    ]

    rows = ws.iter_rows(
        values_only=True
    )

    header = next(rows)

    col = {
        str(v): i
        for i, v in enumerate(header)
    }

    needed = [
        "reference",
        "systematic_name",
        "gene",
        "p_site",
        "p_residue",
        "p_position",
        "treatment_id",
        "fc_log2",
        "p_value",
        "adj_p_value",
    ]

    missing = [
        x
        for x in needed
        if x not in col
    ]

    if missing:
        raise RuntimeError(
            f"Missing Table S6 columns: {missing}"
        )

    out = []

    for row in rows:

        treatment = row[
            col[
                "treatment_id"
            ]
        ]

        if treatment not in {
            HEAT,
            COLD,
        }:
            continue

        out.append({
            k: row[col[k]]
            for k in needed
        })

    atlas_raw = pd.DataFrame(
        out
    )

    CACHE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    atlas_raw.to_csv(
        CACHE,
        sep="\t",
        index=False,
        compression="gzip",
    )


print(
    "Primary PhosphoAtlas rows:",
    len(atlas_raw),
)

print(
    atlas_raw[
        "treatment_id"
    ].value_counts()
)


                                                              
                                            
                                                              

atlas_raw["gene"] = (
    atlas_raw["gene"]
    .where(
        atlas_raw["gene"].notna(),
        "",
    )
    .astype(str)
    .str.strip()
    .str.upper()
)

atlas_raw["p_residue"] = (
    atlas_raw["p_residue"]
    .astype(str)
    .str.strip()
)

atlas_raw["p_position"] = pd.to_numeric(
    atlas_raw["p_position"],
    errors="coerce",
)

atlas_raw["fc_log2"] = pd.to_numeric(
    atlas_raw["fc_log2"],
    errors="coerce",
)

atlas_raw["p_value"] = pd.to_numeric(
    atlas_raw["p_value"],
    errors="coerce",
)

atlas_raw["adj_p_value"] = pd.to_numeric(
    atlas_raw["adj_p_value"],
    errors="coerce",
)


atlas_raw = atlas_raw[
    (atlas_raw["gene"] != "")
    & atlas_raw[
        "p_residue"
    ].isin(
        [
            "S",
            "T",
            "Y",
        ]
    )
    & atlas_raw[
        "p_position"
    ].notna()
    & atlas_raw[
        "fc_log2"
    ].notna()
].copy()


atlas_raw[
    "p_position"
] = atlas_raw[
    "p_position"
].astype(int)


                                                          
atlas_site_treatment = (
    atlas_raw
    .groupby(
        [
            "gene",
            "p_residue",
            "p_position",
            "treatment_id",
        ],
        as_index=False,
    )
    .agg(
        n_rows=(
            "reference",
            "size",
        ),

        fc_log2=(
            "fc_log2",
            "median",
        ),

        p_value=(
            "p_value",
            "min",
        ),

        adj_p_value=(
            "adj_p_value",
            "min",
        ),
    )
)


heat = atlas_site_treatment[
    atlas_site_treatment[
        "treatment_id"
    ] == HEAT
].copy()


cold = atlas_site_treatment[
    atlas_site_treatment[
        "treatment_id"
    ] == COLD
].copy()


heat = heat.rename(
    columns={
        "n_rows":
            "HS42_n_rows",

        "fc_log2":
            "HS42_fc_log2",

        "p_value":
            "HS42_p",

        "adj_p_value":
            "HS42_q",
    }
).drop(
    columns=[
        "treatment_id"
    ]
)


cold = cold.rename(
    columns={
        "n_rows":
            "CS18_n_rows",

        "fc_log2":
            "CS18_fc_log2",

        "p_value":
            "CS18_p",

        "adj_p_value":
            "CS18_q",
    }
).drop(
    columns=[
        "treatment_id"
    ]
)


atlas = heat.merge(
    cold,
    on=[
        "gene",
        "p_residue",
        "p_position",
    ],
    how="inner",
    validate="1:1",
)


atlas[
    "PhosphoAtlas_heat_minus_cold"
] = (
    atlas[
        "HS42_fc_log2"
    ]
    -
    atlas[
        "CS18_fc_log2"
    ]
)


print(
    "Exact phosphosites observed in BOTH HS42 and CS18:",
    len(atlas),
)


                                                              
                                          
                                                              

s1 = pd.read_excel(
    S1FILE,
    sheet_name="Phosphopeptides",
)


required = [
    "id",
    "pSites",
    "Heat.T04",
    "Heat.T06",
    "Cold.T04",
    "Cold.T06",
]

for c in required:

    if c not in s1.columns:
        raise RuntimeError(
            f"Missing Kanshin column {c!r}. "
            f"Available columns include:\n"
            + "\n".join(
                map(
                    str,
                    s1.columns,
                )
            )
        )


gene_col = (
    "Gene"
    if "Gene" in s1.columns
    else "Standard Name"
)


s1[
    "s1_id"
] = s1[
    "id"
].map(
    norm_id
)


s1[
    "gene"
] = (
    s1[
        gene_col
    ]
    .where(
        s1[
            gene_col
        ].notna(),
        "",
    )
    .astype(str)
    .str.strip()
    .str.upper()
)


for c in [
    "Heat.T04",
    "Heat.T06",
    "Cold.T04",
    "Cold.T06",
]:

    s1[c] = pd.to_numeric(
        s1[c],
        errors="coerce",
    )


complete_46 = s1[
    [
        "Heat.T04",
        "Heat.T06",
        "Cold.T04",
        "Cold.T06",
    ]
].notna().all(
    axis=1
)


s1[
    "Kanshin_heat5"
] = np.nan

s1[
    "Kanshin_cold5"
] = np.nan


s1.loc[
    complete_46,
    "Kanshin_heat5",
] = (
    s1.loc[
        complete_46,
        [
            "Heat.T04",
            "Heat.T06",
        ],
    ].mean(
        axis=1
    )
)


s1.loc[
    complete_46,
    "Kanshin_cold5",
] = (
    s1.loc[
        complete_46,
        [
            "Cold.T04",
            "Cold.T06",
        ],
    ].mean(
        axis=1
    )
)


s1[
    "Kanshin_heat_minus_cold_5min"
] = (
    s1[
        "Kanshin_heat5"
    ]
    -
    s1[
        "Kanshin_cold5"
    ]
)


s1[
    "site_id"
] = (
    s1[
        "gene"
    ]
    + "_"
    + s1[
        "pSites"
    ].astype(str)
)


                                                              
                                              
                                                              

atom_rows = []


for _, r in s1.iterrows():

    gene = r[
        "gene"
    ]

    if not gene:
        continue

    sites = parse_sites(
        r[
            "pSites"
        ]
    )

    for site in sites:

        m = re.fullmatch(
            r"([STY])(\d+)",
            site,
        )

        if not m:
            continue

        atom_rows.append({
            "s1_id":
                r[
                    "s1_id"
                ],

            "gene":
                gene,

            "pSites":
                clean(
                    r[
                        "pSites"
                    ]
                ),

            "site_id":
                r[
                    "site_id"
                ],

            "atom_site":
                site,

            "p_residue":
                m.group(1),

            "p_position":
                int(
                    m.group(2)
                ),

            "Kanshin_heat5":
                r[
                    "Kanshin_heat5"
                ],

            "Kanshin_cold5":
                r[
                    "Kanshin_cold5"
                ],

            "Kanshin_heat_minus_cold_5min":
                r[
                    "Kanshin_heat_minus_cold_5min"
                ],
        })


atoms = pd.DataFrame(
    atom_rows
)


mapped_atoms = atoms.merge(
    atlas,
    on=[
        "gene",
        "p_residue",
        "p_position",
    ],
    how="left",
    validate="m:1",
)


mapped_atoms[
    "atlas_match"
] = mapped_atoms[
    "PhosphoAtlas_heat_minus_cold"
].notna()


mapped_atoms.to_csv(
    OUT
    / "ALL_S1_atom_level_PhosphoAtlas_matches.tsv.gz",
    sep="\t",
    index=False,
    compression="gzip",
)


                                                                   
group_rows = []


for (
    s1_id,
    gene,
    psites,
    site_id,
), g in mapped_atoms.groupby(
    [
        "s1_id",
        "gene",
        "pSites",
        "site_id",
    ],
    sort=False,
):

    gm = g[
        g[
            "atlas_match"
        ]
    ].copy()

    if len(gm) == 0:
        continue

    group_rows.append({
        "s1_id":
            s1_id,

        "gene":
            gene,

        "pSites":
            psites,

        "site_id":
            site_id,

        "n_atoms":
            len(g),

        "n_atlas_atoms":
            len(gm),

        "matched_atoms":
            "/".join(
                gm[
                    "atom_site"
                ].astype(str)
            ),

        "Kanshin_heat5":
            g[
                "Kanshin_heat5"
            ].iloc[0],

        "Kanshin_cold5":
            g[
                "Kanshin_cold5"
            ].iloc[0],

        "Kanshin_heat_minus_cold_5min":
            g[
                "Kanshin_heat_minus_cold_5min"
            ].iloc[0],

        "PhosphoAtlas_HS42":
            float(
                np.median(
                    gm[
                        "HS42_fc_log2"
                    ]
                )
            ),

        "PhosphoAtlas_CS18":
            float(
                np.median(
                    gm[
                        "CS18_fc_log2"
                    ]
                )
            ),

        "PhosphoAtlas_heat_minus_cold":
            float(
                np.median(
                    gm[
                        "PhosphoAtlas_heat_minus_cold"
                    ]
                )
            ),

        "HS42_q_min":
            float(
                np.min(
                    gm[
                        "HS42_q"
                    ]
                )
            )
            if gm[
                "HS42_q"
            ].notna().any()
            else np.nan,

        "CS18_q_min":
            float(
                np.min(
                    gm[
                        "CS18_q"
                    ]
                )
            )
            if gm[
                "CS18_q"
            ].notna().any()
            else np.nan,
    })


groups = pd.DataFrame(
    group_rows
)


groups.to_csv(
    OUT
    / "ALL_S1_PhosphoAtlas_group_matches.tsv.gz",
    sep="\t",
    index=False,
    compression="gzip",
)


print(
    "Kanshin S1 groups with exact PhosphoAtlas HS42+CS18 match:",
    len(groups),
)

print(
    "Matched genes:",
    groups[
        "gene"
    ].nunique(),
)


                                                              
                                  
                                                              

direct = pd.read_csv(
    DIRECT_TARGETS,
    sep="\t",
    dtype={
        "s1_id":
            str,
    },
)


direct[
    "s1_id"
] = direct[
    "s1_id"
].map(
    norm_id
)


direct_ids = set(
    direct[
        "s1_id"
    ]
)


if len(
    direct_ids
) != 10:
    raise RuntimeError(
        f"Expected 10 direct-A groups; "
        f"found {len(direct_ids)}"
    )


mem = pd.read_csv(
    NETWORKIN,
    sep="\t",
)


orth = mem[
    mem[
        "analysis_label"
    ] == ORTH_LABEL
].copy()


orth[
    "gene"
] = (
    orth[
        "gene"
    ]
    .astype(str)
    .str.strip()
    .str.upper()
)


orth = orth.drop_duplicates(
    [
        "gene",
        "site_id",
    ]
)


if len(
    orth
) != 20:
    raise RuntimeError(
        f"Expected 20 orthogonal PKA groups; "
        f"found {len(orth)}"
    )


orth_keys = set(
    zip(
        orth[
            "gene"
        ],
        orth[
            "site_id"
        ],
    )
)


orth_s1 = s1[
    s1.apply(
        lambda r: (
            r[
                "gene"
            ],
            r[
                "site_id"
            ],
        ) in orth_keys,
        axis=1,
    )
].copy()


if len(
    orth_s1
) != 20:
    print(
        "ORTHOGONAL MEMBERSHIP:"
    )

    print(
        orth[
            [
                "gene",
                "site_id",
            ]
        ].to_string(
            index=False
        )
    )

    print(
        "\nS1 MATCHES:"
    )

    print(
        orth_s1[
            [
                "s1_id",
                "gene",
                "site_id",
            ]
        ].to_string(
            index=False
        )
    )

    raise RuntimeError(
        f"Expected 20 S1 matches for orthogonal PKA; "
        f"found {len(orth_s1)}"
    )


orth_ids = set(
    orth_s1[
        "s1_id"
    ]
)


                                                              
                                          
                                                              

def module_test(
    name,
    target_ids,
    seed,
):

    target_all = s1[
        s1[
            "s1_id"
        ].isin(
            target_ids
        )
    ].copy()

    target = groups[
        groups[
            "s1_id"
        ].isin(
            target_ids
        )
    ].copy()


    target_genes = set(
        target[
            "gene"
        ]
    )


                                     
    bg = groups[
        ~groups[
            "gene"
        ].isin(
            target_genes
        )
        &
        ~groups[
            "s1_id"
        ].isin(
            target_ids
        )
    ].copy()


    target_gene = (
        target
        .groupby(
            "gene",
            as_index=False,
        )
        .agg(
            n_groups=(
                "s1_id",
                "nunique",
            ),

            atlas_delta=(
                "PhosphoAtlas_heat_minus_cold",
                "median",
            ),
        )
    )


    if len(
        target_gene
    ) < 3:
        raise RuntimeError(
            f"{name}: only "
            f"{len(target_gene)} externally matched genes"
        )


    mults = sorted(
        target_gene[
            "n_groups"
        ].astype(int).tolist(),
        reverse=True,
    )


    observed = float(
        target_gene[
            "atlas_delta"
        ].mean()
    )


    observed_median = float(
        target_gene[
            "atlas_delta"
        ].median()
    )


    negative_genes = int(
        (
            target_gene[
                "atlas_delta"
            ] < 0
        ).sum()
    )


    bg_by_gene = {
        gene:
            g[
                "PhosphoAtlas_heat_minus_cold"
            ].to_numpy(
                dtype=float
            )

        for gene, g in bg.groupby(
            "gene"
        )
    }


    eligible = {
        m: np.array(
            [
                gene
                for gene, vals
                in bg_by_gene.items()
                if len(vals) >= m
            ],
            dtype=object,
        )

        for m in set(
            mults
        )
    }


    for m in sorted(
        eligible
    ):

        print(
            f"{name}: background genes "
            f"with >= {m} matched groups:",
            len(
                eligible[m]
            ),
        )

        if len(
            eligible[m]
        ) < len(
            mults
        ):
            raise RuntimeError(
                f"{name}: insufficient null genes "
                f"for multiplicity {m}"
            )


    rng = np.random.default_rng(
        seed
    )

    null_mean = np.empty(
        N_PERM,
        dtype=float,
    )

    null_median = np.empty(
        N_PERM,
        dtype=float,
    )

    null_negative = np.empty(
        N_PERM,
        dtype=int,
    )


    for i in range(
        N_PERM
    ):

        chosen = set()
        values = []

        for m in mults:

            pool = eligible[m]

            while True:

                gene = pool[
                    rng.integers(
                        0,
                        len(pool),
                    )
                ]

                if gene not in chosen:
                    break

            chosen.add(
                gene
            )

            vals = bg_by_gene[
                gene
            ]

            idx = rng.choice(
                len(vals),
                size=m,
                replace=False,
            )

            values.append(
                float(
                    np.median(
                        vals[
                            idx
                        ]
                    )
                )
            )


        values = np.asarray(
            values
        )

        null_mean[i] = np.mean(
            values
        )

        null_median[i] = np.median(
            values
        )

        null_negative[i] = np.sum(
            values < 0
        )


    null_mu = float(
        np.mean(
            null_mean
        )
    )

    null_sd = float(
        np.std(
            null_mean,
            ddof=1,
        )
    )


    p_mean = empirical_lower(
        observed,
        null_mean,
    )


    p_median = empirical_lower(
        observed_median,
        null_median,
    )


    p_sign = float(
        (
            1
            + np.sum(
                null_negative
                >= negative_genes
            )
        )
        / (
            N_PERM
            + 1
        )
    )


                                                              
                                                     
                                                              

    concord = target[
        target[
            "Kanshin_heat_minus_cold_5min"
        ].notna()
        &
        target[
            "PhosphoAtlas_heat_minus_cold"
        ].notna()
    ].copy()


    gene_concord = (
        concord
        .groupby(
            "gene",
            as_index=False,
        )
        .agg(
            Kanshin_delta5=(
                "Kanshin_heat_minus_cold_5min",
                "median",
            ),

            Atlas_delta5=(
                "PhosphoAtlas_heat_minus_cold",
                "median",
            ),
        )
    )


    if len(
        gene_concord
    ) >= 3:

        spear = spearmanr(
            gene_concord[
                "Kanshin_delta5"
            ],
            gene_concord[
                "Atlas_delta5"
            ],
        )

        pear = pearsonr(
            gene_concord[
                "Kanshin_delta5"
            ],
            gene_concord[
                "Atlas_delta5"
            ],
        )

        same_sign = (
            np.sign(
                gene_concord[
                    "Kanshin_delta5"
                ]
            )
            ==
            np.sign(
                gene_concord[
                    "Atlas_delta5"
                ]
            )
        )

        n_same = int(
            same_sign.sum()
        )

        n_conc = len(
            gene_concord
        )

        sign_p = float(
            binomtest(
                n_same,
                n_conc,
                p=0.5,
                alternative="greater",
            ).pvalue
        )

        spear_r = float(
            spear.statistic
        )

        spear_p = float(
            spear.pvalue
        )

        pear_r = float(
            pear.statistic
        )

        pear_p = float(
            pear.pvalue
        )

    else:

        n_same = np.nan
        n_conc = len(
            gene_concord
        )

        sign_p = np.nan
        spear_r = np.nan
        spear_p = np.nan
        pear_r = np.nan
        pear_p = np.nan


    target.to_csv(
        OUT
        / f"{name}_matched_groups.tsv",
        sep="\t",
        index=False,
    )


    target_gene.to_csv(
        OUT
        / f"{name}_external_gene_effects.tsv",
        sep="\t",
        index=False,
    )


    gene_concord.to_csv(
        OUT
        / f"{name}_Kanshin_vs_Atlas_gene_concordance.tsv",
        sep="\t",
        index=False,
    )


    result = {
        "module":
            name,

        "predefined_groups":
            len(
                target_ids
            ),

        "predefined_genes":
            target_all[
                "gene"
            ].nunique(),

        "atlas_matched_groups":
            len(
                target
            ),

        "atlas_matched_genes":
            len(
                target_gene
            ),

        "multiplicity_pattern":
            ",".join(
                map(
                    str,
                    mults,
                )
            ),

        "observed_mean_Atlas_HS42_minus_CS18":
            observed,

        "observed_median_Atlas_HS42_minus_CS18":
            observed_median,

        "negative_genes":
            negative_genes,

        "null_mean":
            null_mu,

        "null_sd":
            null_sd,

        "empirical_z":
            float(
                (
                    observed
                    - null_mu
                )
                / null_sd
            ),

        "mean_empirical_lower_p":
            p_mean,

        "median_empirical_lower_p":
            p_median,

        "negative_count_empirical_p":
            p_sign,

        "n_genes_with_complete_Kanshin_4_6min":
            n_conc,

        "same_sign_genes":
            n_same,

        "sign_concordance_binomial_p":
            sign_p,

        "spearman_r":
            spear_r,

        "spearman_p":
            spear_p,

        "pearson_r":
            pear_r,

        "pearson_p":
            pear_p,
    }


    return result


                                                              
                                           
                                                              

results = []


results.append(
    module_test(
        "DIRECT_A_PKA",
        direct_ids,
        SEED + 1,
    )
)


results.append(
    module_test(
        "ORTHOGONAL_NETWORKIN_PKA",
        orth_ids,
        SEED + 2,
    )
)


results = pd.DataFrame(
    results
)


results[
    "mean_empirical_q_across_two_modules"
] = bh(
    results[
        "mean_empirical_lower_p"
    ]
)


results.to_csv(
    OUT
    / "PHOSPHOATLAS_PRIMARY_REPLICATION_SUMMARY.tsv",
    sep="\t",
    index=False,
)


                                                              
          
                                                              

print()
print("=" * 120)
print("PHOSPHOATLAS 42C-vs-18C 5-MIN INDEPENDENT REPLICATION")
print("=" * 120)

print()
print(
    results.to_string(
        index=False
    )
)


for name in [
    "DIRECT_A_PKA",
    "ORTHOGONAL_NETWORKIN_PKA",
]:

    print()
    print("=" * 120)
    print(name, "- MATCHED GROUPS")
    print("=" * 120)

    d = pd.read_csv(
        OUT
        / f"{name}_matched_groups.tsv",
        sep="\t",
    )

    cols = [
        "gene",
        "pSites",
        "matched_atoms",
        "Kanshin_heat_minus_cold_5min",
        "PhosphoAtlas_HS42",
        "PhosphoAtlas_CS18",
        "PhosphoAtlas_heat_minus_cold",
        "HS42_q_min",
        "CS18_q_min",
    ]

    print(
        d[
            cols
        ].to_string(
            index=False
        )
    )


print()
print("Output:", OUT)
